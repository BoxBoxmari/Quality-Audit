"""
Pytest configuration and shared fixtures.
"""

import pytest
from docx import Document


@pytest.fixture
def sample_word_file(tmp_path):
    """
    Create a sample Word document for testing.

    Returns:
        str: Path to the created .docx file
    """
    doc_path = tmp_path / "sample.docx"
    doc = Document()

    # Add a simple table for testing
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Account"
    table.cell(0, 1).text = "2024"
    table.cell(0, 2).text = "2023"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "1000"
    table.cell(1, 2).text = "900"
    table.cell(2, 0).text = "Total"
    table.cell(2, 1).text = "1000"
    table.cell(2, 2).text = "900"

    doc.save(str(doc_path))
    return str(doc_path)
